from openpyxl import load_workbook
from skimage import color
from PIL import Image
import numpy as np

INPUT_XLSX = "telemetry.xlsx"
OUTPUT_PNG = "reconstructed.png"

print("Loading workbook...")

wb = load_workbook(INPUT_XLSX, data_only=True)

sheet_names = wb.sheetnames[:3]

print("Sheets:", sheet_names)

channels = []

for name in sheet_names:
    ws = wb[name]

    rows = []

    for row in ws.iter_rows(values_only=True):
        rows.append([
            float(c) if c is not None else 0.0
            for c in row
        ])

    arr = np.array(rows, dtype=float)

    print(
        f"{name}: shape={arr.shape}, "
        f"min={arr.min():.3f}, "
        f"max={arr.max():.3f}"
    )

    channels.append(arr)


lab = np.stack(channels, axis=-1)

print("Converting CIELAB to sRGB...")

rgb = color.lab2rgb(lab)

# ROUND to nearest integer.
rgb8 = np.rint(
    np.clip(rgb, 0, 1) * 255
).astype(np.uint8)

img = Image.fromarray(rgb8, "RGB")
img.save(OUTPUT_PNG)

print(f"Saved reconstructed image: {OUTPUT_PNG}")
print("Image size:", img.size)